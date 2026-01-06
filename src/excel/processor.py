"""
Excel处理模块
负责读取、解析和分析Excel文件
"""
import os
from typing import Dict, List, Any, Optional, Tuple
from dataclasses import dataclass
import pandas as pd
import openpyxl
from openpyxl.utils import get_column_letter
from loguru import logger


@dataclass
class CellInfo:
    """单元格信息"""
    row: int
    col: int
    value: Any
    formula: Optional[str] = None
    data_type: str = "string"


@dataclass
class SheetInfo:
    """工作表信息"""
    name: str
    index: int
    rows: int
    cols: int
    has_merged_cells: bool
    merged_ranges: List[Tuple[int, int, int, int]]  # (min_row, min_col, max_row, max_col)
    headers: List[str]
    header_row: Optional[int]


@dataclass
class ExcelInfo:
    """Excel文件信息"""
    file_path: str
    file_name: str
    file_size: int
    sheets: List[SheetInfo]
    total_sheets: int
    metadata: Dict[str, Any]


class ExcelProcessor:
    """Excel处理器"""
    
    def __init__(self, max_rows: int = 10000, max_cols: int = 500):
        """
        初始化Excel处理器
        
        Args:
            max_rows: 最大行数限制
            max_cols: 最大列数限制
        """
        self.max_rows = max_rows
        self.max_cols = max_cols
        logger.info(f"Excel处理器初始化完成，最大行数: {max_rows}, 最大列数: {max_cols}")
    
    def load_excel(self, file_path: str) -> ExcelInfo:
        """
        加载Excel文件
        
        Args:
            file_path: Excel文件路径
            
        Returns:
            ExcelInfo: Excel文件信息对象
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"文件不存在: {file_path}")
        
        file_size = os.path.getsize(file_path)
        file_name = os.path.basename(file_path)
        
        logger.info(f"开始加载Excel文件: {file_name} (大小: {file_size} bytes)")
        
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
            sheets = []
            
            for idx, sheet_name in enumerate(wb.sheetnames):
                sheet = wb[sheet_name]
                sheet_info = self._analyze_sheet(sheet, idx)
                sheets.append(sheet_info)
                logger.info(f"工作表 {idx+1}/{len(wb.sheetnames)}: {sheet_name} ({sheet_info.rows}行 x {sheet_info.cols}列)")
            
            wb.close()
            
            excel_info = ExcelInfo(
                file_path=file_path,
                file_name=file_name,
                file_size=file_size,
                sheets=sheets,
                total_sheets=len(sheets),
                metadata={
                    "format": os.path.splitext(file_name)[1],
                    "created_time": os.path.getctime(file_path),
                    "modified_time": os.path.getmtime(file_path)
                }
            )
            
            logger.info(f"Excel文件加载完成: {file_name}")
            return excel_info
            
        except Exception as e:
            logger.error(f"加载Excel文件失败: {e}")
            raise
    
    def _analyze_sheet(self, sheet, index: int) -> SheetInfo:
        """
        分析工作表结构
        
        Args:
            sheet: openpyxl工作表对象
            index: 工作表索引
            
        Returns:
            SheetInfo: 工作表信息
        """
        # 获取工作表维度
        max_row = min(sheet.max_row, self.max_rows)
        max_col = min(sheet.max_column, self.max_cols)
        
        # 检查合并单元格
        merged_ranges = []
        try:
            for merged_range in sheet.merged_cells.ranges:
                min_row, min_col, max_row_m, max_col_m = merged_range.bounds
                merged_ranges.append((min_row, min_col, max_row_m, max_col_m))
        except AttributeError:
            # ReadOnlyWorksheet不支持merged_cells，跳过
            pass
        
        has_merged = len(merged_ranges) > 0
        
        # 识别表头
        headers, header_row = self._identify_headers(sheet, max_row, max_col)
        
        return SheetInfo(
            name=sheet.title,
            index=index,
            rows=max_row,
            cols=max_col,
            has_merged_cells=has_merged,
            merged_ranges=merged_ranges,
            headers=headers,
            header_row=header_row
        )
    
    def _identify_headers(self, sheet, max_row: int, max_col: int) -> Tuple[List[str], Optional[int]]:
        """
        识别表头
        
        Args:
            sheet: openpyxl工作表对象
            max_row: 最大行数
            max_col: 最大列数
            
        Returns:
            Tuple[List[str], Optional[int]]: (表头列表, 表头行号)
        """
        # 简单策略：检查前3行，找到非空单元格最多的一行
        best_row = None
        best_headers = []
        max_non_empty = 0
        
        for row_idx in range(1, min(4, max_row + 1)):
            headers = []
            non_empty_count = 0
            
            for col_idx in range(1, max_col + 1):
                cell_value = sheet.cell(row=row_idx, column=col_idx).value
                if cell_value is not None:
                    headers.append(str(cell_value))
                    non_empty_count += 1
                else:
                    headers.append("")
            
            if non_empty_count > max_non_empty:
                max_non_empty = non_empty_count
                best_headers = headers
                best_row = row_idx
        
        return best_headers, best_row
    
    def read_sheet_data(self, file_path: str, sheet_name: str, 
                       header_row: int = 0) -> pd.DataFrame:
        """
        读取工作表数据为DataFrame
        
        Args:
            file_path: Excel文件路径
            sheet_name: 工作表名称
            header_row: 表头行号（0-based）
            
        Returns:
            pd.DataFrame: 数据框
        """
        try:
            df = pd.read_excel(
                file_path,
                sheet_name=sheet_name,
                header=header_row,
                nrows=self.max_rows,
                usecols=lambda x: x < self.max_cols
            )
            
            logger.info(f"读取工作表数据成功: {sheet_name}, 形状: {df.shape}")
            return df
            
        except Exception as e:
            logger.error(f"读取工作表数据失败: {e}")
            raise
    
    def extract_cell_values(self, file_path: str, sheet_name: str, 
                           cell_range: str) -> List[CellInfo]:
        """
        提取指定范围内单元格的值
        
        Args:
            file_path: Excel文件路径
            sheet_name: 工作表名称
            cell_range: 单元格范围，如 "A1:C10"
            
        Returns:
            List[CellInfo]: 单元格信息列表
        """
        try:
            wb = openpyxl.load_workbook(file_path, data_only=False)
            sheet = wb[sheet_name]
            
            cells = []
            for row in sheet[cell_range]:
                for cell in row:
                    cells.append(CellInfo(
                        row=cell.row,
                        col=cell.column,
                        value=cell.value,
                        formula=cell.data_type == 'f' and cell.value or None,
                        data_type=str(cell.data_type)
                    ))
            
            wb.close()
            logger.info(f"提取单元格值成功: {sheet_name}, 范围: {cell_range}, 数量: {len(cells)}")
            return cells
            
        except Exception as e:
            logger.error(f"提取单元格值失败: {e}")
            raise
    
    def get_sheet_data_sample(self, file_path: str, sheet_name: str, 
                             n_rows: int = 10) -> Dict[str, Any]:
        """
        获取工作表数据样本
        
        Args:
            file_path: Excel文件路径
            sheet_name: 工作表名称
            n_rows: 采样行数
            
        Returns:
            Dict[str, Any]: 包含样本数据的字典
        """
        try:
            df = pd.read_excel(
                file_path,
                sheet_name=sheet_name,
                nrows=n_rows,
                header=None
            )
            
            # 尝试识别表头
            header_row = self._guess_header_row(df)
            if header_row is not None:
                df = pd.read_excel(
                    file_path,
                    sheet_name=sheet_name,
                    header=header_row,
                    nrows=n_rows + 1
                )
            
            sample = {
                "sheet_name": sheet_name,
                "header_row": header_row,
                "columns": df.columns.tolist(),
                "data": df.head(n_rows).to_dict(orient='records'),
                "shape": df.shape,
                "dtypes": df.dtypes.astype(str).to_dict()
            }
            
            logger.info(f"获取数据样本成功: {sheet_name}")
            return sample
            
        except Exception as e:
            logger.error(f"获取数据样本失败: {e}")
            raise
    
    def _guess_header_row(self, df: pd.DataFrame) -> Optional[int]:
        """
        猜测表头行号
        
        Args:
            df: 数据框
            
        Returns:
            Optional[int]: 表头行号
        """
        for idx in range(min(3, len(df))):
            row = df.iloc[idx]
            # 检查是否有足够的非空值
            non_empty = sum(1 for v in row if v is not None and str(v).strip() != "")
            if non_empty >= len(row) * 0.5:
                return idx
        return None
    
    def validate_excel(self, file_path: str) -> Tuple[bool, List[str]]:
        """
        验证Excel文件
        
        Args:
            file_path: Excel文件路径
            
        Returns:
            Tuple[bool, List[str]]: (是否有效, 错误/警告信息列表)
        """
        errors = []
        
        if not os.path.exists(file_path):
            errors.append(f"文件不存在: {file_path}")
            return False, errors
        
        file_ext = os.path.splitext(file_path)[1].lower()
        if file_ext not in ['.xlsx', '.xls', '.xlsm']:
            errors.append(f"不支持的文件格式: {file_ext}")
            return False, errors
        
        try:
            wb = openpyxl.load_workbook(file_path, read_only=True)
            
            if len(wb.sheetnames) == 0:
                errors.append("Excel文件中没有工作表")
            
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                if sheet.max_row > self.max_rows:
                    errors.append(f"工作表 '{sheet_name}' 行数超过限制: {sheet.max_row} > {self.max_rows}")
                if sheet.max_column > self.max_cols:
                    errors.append(f"工作表 '{sheet_name}' 列数超过限制: {sheet.max_column} > {self.max_cols}")
            
            wb.close()
            
        except Exception as e:
            errors.append(f"Excel文件解析失败: {str(e)}")
            return False, errors
        
        return len(errors) == 0, errors